#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
AudioLab MZ - Universal Document Extraction Engine v2.5 (Clean Architecture)
High-modularity, Strategy Pattern extraction pipeline.
Supports:
  - PDF: PyMuPDF with OCR fallback via Tesseract
  - DOCX: python-docx with paragraph and table extraction
  - TXT / Markdown: Clean UTF-8 reading and header segmentation
  - Images: Tesseract OCR with PIL image pre-processing
  - Audio: ffmpeg 16kHz WAV chunking + SpeechRecognition STT
"""

import sys
import os
import json
import re
import argparse
import tempfile
import glob
import subprocess
import shutil
import struct
from dataclasses import dataclass, field
from typing import List, Dict, Optional, Tuple, Type

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')
if hasattr(sys.stderr, 'reconfigure'):
    sys.stderr.reconfigure(encoding='utf-8')


# ──────────────────────────────────────────────────────────────────────────────
# 1. TEXT NORMALIZER (Single Purpose: Sanitizing & Cleaning Text)
# ──────────────────────────────────────────────────────────────────────────────

class TextNormalizer:
    """Handles character normalization, whitespace fixing, and OCR artifact cleaning."""

    @staticmethod
    def clean(text: Optional[str]) -> str:
        if not text:
            return ""

        text = text.replace('\r\n', '\n').replace('\r', '\n')

        # Fix hyphenated words at line breaks (e.g., transfor-\nmación → transformación)
        text = re.sub(r'([a-zA-ZáéíóúÁÉÍÓÚñÑ]+)-\n([a-zA-ZáéíóúÁÉÍÓÚñÑ]+)', r'\1\2', text)

        # Fix soft hyphens and unusual whitespace
        text = text.replace('\xad', '').replace('\xa0', ' ')
        text = re.sub(r'\[(?:pic|image|imagen)\]', '', text, flags=re.IGNORECASE)
        text = re.sub(r'[ \t\f\v]+', ' ', text)

        # Remove standalone page numbers
        text = re.sub(
            r'^\s*(?:Página|Pág\.?|Page)?\s*\d+\s*(?:de|\/)\s*\d+\s*$',
            '',
            text,
            flags=re.MULTILINE | re.IGNORECASE
        )
        text = re.sub(r'^\s*\d+\s*$', '', text, flags=re.MULTILINE)

        # Remove ASCII control characters
        text = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]', '', text)

        # Consolidate multiple linebreaks
        text = re.sub(r'\n{3,}', '\n\n', text)

        # Bridge single linebreaks into a space for NLP continuity
        text = re.sub(r'(?<!\n)\n(?!\n)', ' ', text)

        # Consolidate consecutive spaces
        text = re.sub(r'[ \t]+', ' ', text)

        return text.strip()


# ──────────────────────────────────────────────────────────────────────────────
# 2. DATA TRANSFER OBJECT (DTO)
# ──────────────────────────────────────────────────────────────────────────────

@dataclass
class ExtractionResult:
    title: str = ""
    author: str = ""
    pages_text: List[str] = field(default_factory=list)
    ocr_used: bool = False


# ──────────────────────────────────────────────────────────────────────────────
# 3. EXTRACTOR STRATEGIES (Single Responsibility per Document Type)
# ──────────────────────────────────────────────────────────────────────────────

class BaseExtractor:
    """Abstract base strategy for document text extractors."""
    def extract(self, file_path: str) -> ExtractionResult:
        raise NotImplementedError


class PdfExtractor(BaseExtractor):
    """Extracts text from PDF documents using PyMuPDF, with automatic Tesseract OCR fallback for scanned pages."""

    SCANNED_CHAR_THRESHOLD = 50

    def extract(self, file_path: str) -> ExtractionResult:
        if self._is_scanned(file_path):
            sys.stderr.write("[INFO] PDF escaneado detectado — activando OCR con Tesseract...\n")
            return self._extract_with_ocr(file_path)
        return self._extract_digital(file_path)

    def _is_scanned(self, file_path: str) -> bool:
        try:
            import pymupdf as fitz
            doc = fitz.open(file_path)
            total_chars = sum(len(page.get_text("text").strip()) for page in doc)
            doc.close()
            return total_chars < self.SCANNED_CHAR_THRESHOLD
        except Exception:
            return False

    def _extract_digital(self, file_path: str) -> ExtractionResult:
        try:
            import pymupdf as fitz
            doc = fitz.open(file_path)
            metadata = doc.metadata or {}
            title = metadata.get('title', '').strip()
            author = metadata.get('author', '').strip()

            pages_text = []
            for page in doc:
                cleaned = TextNormalizer.clean(page.get_text("text"))
                if cleaned:
                    pages_text.append(cleaned)
            doc.close()
            return ExtractionResult(title=title, author=author, pages_text=pages_text, ocr_used=False)
        except Exception as e1:
            try:
                import pypdf
                reader = pypdf.PdfReader(file_path)
                title = reader.metadata.get('/Title', '') if reader.metadata else ''
                author = reader.metadata.get('/Author', '') if reader.metadata else ''
                pages_text = [
                    TextNormalizer.clean(page.extract_text() or '')
                    for page in reader.pages
                    if TextNormalizer.clean(page.extract_text() or '')
                ]
                return ExtractionResult(title=title or '', author=author or '', pages_text=pages_text, ocr_used=False)
            except Exception as e2:
                raise Exception(f"Error procesando PDF digital: {e1} / {e2}")

    def _extract_with_ocr(self, file_path: str) -> ExtractionResult:
        try:
            import pymupdf as fitz
            import pytesseract
            from PIL import Image
            import io
        except ImportError as e:
            raise Exception(f"Tesseract OCR requerido para PDF escaneado: {e}")

        doc = fitz.open(file_path)
        metadata = doc.metadata or {}
        title = metadata.get('title', '').strip()
        author = metadata.get('author', '').strip()

        pages_text = []
        matrix = fitz.Matrix(200 / 72, 200 / 72)

        for page_num, page in enumerate(doc):
            try:
                pix = page.get_pixmap(matrix=matrix)
                img = Image.open(io.BytesIO(pix.tobytes("png")))
                raw_text = pytesseract.image_to_string(img, lang='spa+eng', config=r'--oem 3 --psm 3')
                cleaned = TextNormalizer.clean(raw_text)
                if cleaned:
                    pages_text.append(cleaned)
            except Exception as e:
                sys.stderr.write(f"[OCR] Página {page_num + 1} omitida: {e}\n")
                continue

        doc.close()
        return ExtractionResult(title=title, author=author, pages_text=pages_text, ocr_used=True)


class DocxExtractor(BaseExtractor):
    """Extracts text and table data from Microsoft Word (.docx) documents."""

    def extract(self, file_path: str) -> ExtractionResult:
        from docx import Document
        doc = Document(file_path)

        title = ""
        author = ""
        try:
            props = doc.core_properties
            title = props.title or ""
            author = props.author or ""
        except Exception:
            pass

        paragraphs = []
        current_section = []

        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue
            style_name = para.style.name if para.style else ""
            if style_name.startswith(('Heading', 'Título', 'Title', 'Subtitle')):
                if current_section:
                    paragraphs.append('\n\n'.join(current_section))
                    current_section = []
                current_section.append(f"\n{text}\n")
            else:
                current_section.append(text)

        for table in doc.tables:
            table_rows = []
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    deduped = [cells[0]]
                    for c in cells[1:]:
                        if c != deduped[-1]:
                            deduped.append(c)
                    table_rows.append(' | '.join(deduped))
            if table_rows:
                paragraphs.append('\n'.join(table_rows))

        if current_section:
            paragraphs.append('\n\n'.join(current_section))

        pages_text = [TextNormalizer.clean(p) for p in paragraphs if TextNormalizer.clean(p)]
        return ExtractionResult(title=title, author=author, pages_text=pages_text, ocr_used=False)


class DocExtractor(BaseExtractor):
    """Extracts text from binary Microsoft Word 97-2003 (.doc) documents using antiword or built-in OLE2 parser."""

    def _normalize_masonic_glyphs(self, text: str) -> str:
        """Reconstructs authentic tripunctuated symbols (∴) when legacy Word documents used custom fonts mapping to '('."""
        return re.sub(
            r'\b(A|L|G|D|U|I|Ven|Q|QQ|H|HH|VVig|Vig|VVisit|Visit|S|F|E|V|M|MM|Or|Resp|Log|Secr|Orad|Tes|Hosp|Exp|Prof)\s*\(\s*',
            r'\1∴ ',
            text,
            flags=re.IGNORECASE
        )

    def extract(self, file_path: str) -> ExtractionResult:
        # 1. Try antiword if available (standard in Linux container)
        antiword_cmd = shutil.which('antiword') or ('antiword' if os.name != 'nt' else None)
        if antiword_cmd:
            try:
                proc = subprocess.run(
                    [antiword_cmd, '-m', 'UTF-8', file_path],
                    capture_output=True,
                    text=True,
                    check=True
                )
                if proc.stdout.strip():
                    parts = proc.stdout.split('\n\n')
                    pages_text = [TextNormalizer.clean(self._normalize_masonic_glyphs(p)) for p in parts if TextNormalizer.clean(p)]
                    if pages_text:
                        return ExtractionResult(title='', author='', pages_text=pages_text, ocr_used=False)
            except Exception as e:
                sys.stderr.write(f"[DOC] antiword fallo, intentando extractor OLE2 nativo: {e}\n")

        # 2. Resilient Pure-Python OLE2 Piece Table Parser
        try:
            with open(file_path, 'rb') as f:
                data = f.read()

            if len(data) >= 512 and data[:8] == b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1':
                fat_sec = struct.unpack_from('<I', data, 76)[0]
                fat_offset = (fat_sec + 1) * 512
                fat = struct.unpack_from(f'<{128}I', data, fat_offset)

                def get_stream(start_sec, size):
                    cur = start_sec
                    res = bytearray()
                    while cur < 0xFFFFFFFE and len(res) < size:
                        offset = (cur + 1) * 512
                        chunk = data[offset : offset + 512]
                        res.extend(chunk)
                        cur = fat[cur]
                    return bytes(res[:size])

                sec_dir = struct.unpack_from('<I', data, 48)[0]
                dir_offset = (sec_dir + 1) * 512
                streams = {}
                for i in range(32):
                    entry = data[dir_offset + i*128 : dir_offset + (i+1)*128]
                    if len(entry) < 128:
                        break
                    name_len = struct.unpack_from('<H', entry, 64)[0]
                    if name_len == 0:
                        continue
                    name = entry[:name_len-2].decode('utf-16le', errors='ignore')
                    sec_start = struct.unpack_from('<I', entry, 116)[0]
                    size = struct.unpack_from('<I', entry, 120)[0]
                    streams[name] = (sec_start, size)

                if 'WordDocument' in streams:
                    word_doc = get_stream(*streams['WordDocument'])
                    flags = struct.unpack_from('<H', word_doc, 0x000A)[0]
                    tbl_name = '1Table' if (flags & 0x0200) else '0Table'
                    if tbl_name in streams:
                        tbl_doc = get_stream(*streams[tbl_name])
                        fcClx = struct.unpack_from('<I', word_doc, 0x01A2)[0]
                        lcbClx = struct.unpack_from('<I', word_doc, 0x01A6)[0]
                        if fcClx < len(tbl_doc):
                            clx = tbl_doc[fcClx : fcClx + lcbClx]
                            if len(clx) >= 5 and clx[0] == 2:
                                cb = struct.unpack_from('<I', clx, 1)[0]
                                n = (cb - 4) // 12
                                cps = struct.unpack_from(f'<{n+1}I', clx, 5)
                                pcds_offset = 5 + (n+1)*4
                                text_pieces = []
                                for i in range(n):
                                    fc = struct.unpack_from('<I', clx, pcds_offset + i*8 + 2)[0]
                                    is_unicode = (fc & 0x40000000) == 0
                                    raw_fc = fc & 0x3FFFFFFF
                                    cp_len = cps[i+1] - cps[i]
                                    if not is_unicode:
                                        offset = raw_fc // 2
                                        text_pieces.append(word_doc[offset : offset + cp_len].decode('cp1252', errors='replace'))
                                    else:
                                        offset = raw_fc
                                        text_pieces.append(word_doc[offset : offset + cp_len*2].decode('utf-16le', errors='replace'))

                                full_text = ''.join(text_pieces)
                                parts = full_text.split('\r')
                                pages_text = [TextNormalizer.clean(self._normalize_masonic_glyphs(p)) for p in parts if TextNormalizer.clean(p)]
                                if pages_text:
                                    return ExtractionResult(title='', author='', pages_text=pages_text, ocr_used=False)
        except Exception as e:
            sys.stderr.write(f"[DOC] Parser OLE2 fallo: {e}\n")

        # 3. Last-resort fallback: extract printable character sequences
        try:
            with open(file_path, 'rb') as f:
                raw_bytes = f.read()
            matches = re.findall(b'[\x20-\x7e\xa0-\xff]{10,}', raw_bytes)
            lines = [m.decode('cp1252', errors='ignore') for m in matches]
            pages_text = [TextNormalizer.clean(p) for p in lines if TextNormalizer.clean(p)]
            return ExtractionResult(title='', author='', pages_text=pages_text, ocr_used=False)
        except Exception as e:
            raise Exception(f"No se pudo extraer texto del archivo .doc: {e}")


class PlainTextExtractor(BaseExtractor):
    """Extracts text from plaintext (.txt) and Markdown (.md) documents."""

    def extract(self, file_path: str) -> ExtractionResult:
        is_markdown = file_path.lower().endswith(('.md', '.markdown'))

        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                raw = f.read()
        except UnicodeDecodeError:
            with open(file_path, 'r', encoding='latin-1') as f:
                raw = f.read()

        title = ""
        author = ""

        if is_markdown:
            h1_match = re.search(r'^#\s+(.+)$', raw, re.MULTILINE)
            if h1_match:
                title = h1_match.group(1).strip()
            parts = re.split(r'(?=^#{1,3}\s+)', raw, flags=re.MULTILINE)
        else:
            parts = raw.split('\n\n')

        pages_text = [TextNormalizer.clean(p) for p in parts if TextNormalizer.clean(p)]
        return ExtractionResult(title=title, author=author, pages_text=pages_text, ocr_used=False)


class ImageOcrExtractor(BaseExtractor):
    """Performs Optical Character Recognition (OCR) on image files (.png, .jpg, .webp, etc.)."""

    def extract(self, file_path: str) -> ExtractionResult:
        try:
            import pytesseract
            from PIL import Image
        except ImportError:
            raise Exception("El motor OCR (pytesseract/Pillow) no está disponible en el entorno de ejecución.")

        try:
            img = Image.open(file_path)
            if img.mode not in ('L', 'RGB'):
                img = img.convert('RGB')

            custom_config = r'--oem 3 --psm 3'
            try:
                raw_text = pytesseract.image_to_string(img, lang='spa+eng', config=custom_config)
            except pytesseract.TesseractNotFoundError:
                raise Exception("El binario de Tesseract OCR no está instalado en el servidor.")

            cleaned = TextNormalizer.clean(raw_text)

            if not cleaned or len(cleaned.split()) < 3:
                raw_text_fb = pytesseract.image_to_string(img, lang='spa+eng', config=r'--oem 3 --psm 6')
                cleaned_fb = TextNormalizer.clean(raw_text_fb)
                if cleaned_fb:
                    cleaned = cleaned_fb

            pages_text = [cleaned] if cleaned else []
            return ExtractionResult(title="", author="OCR de Imagen", pages_text=pages_text, ocr_used=True)
        except Exception as e:
            if "Tesseract" in str(e) or "motor OCR" in str(e):
                raise
            raise Exception(f"Error procesando OCR de la imagen: {e}")


class AudioTranscriptionExtractor(BaseExtractor):
    """Converts spoken audio files (.mp3, .wav, .m4a, .ogg) to text using ffmpeg and SpeechRecognition."""

    CHUNK_DURATION_SECONDS = "45"

    @staticmethod
    def _get_ffmpeg_binary() -> str:
        found = shutil.which("ffmpeg")
        if found:
            return found
        candidates = [
            r"C:\Program Files\BlueStacks_nxt\ffmpeg.exe",
            r"C:\Program Files\ZWSOFT\ZW3D 2026\ffmpeg.exe",
            "/usr/bin/ffmpeg",
            "/usr/local/bin/ffmpeg",
        ]
        for c in candidates:
            if os.path.exists(c):
                return c
        return "ffmpeg"

    def extract(self, file_path: str) -> ExtractionResult:
        import subprocess
        try:
            import speech_recognition as sr
        except ImportError:
            raise Exception("El módulo SpeechRecognition no está instalado en el entorno de Python.")

        ffmpeg_bin = self._get_ffmpeg_binary()
        try:
            subprocess.run([ffmpeg_bin, "-version"], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, check=True)
        except Exception:
            raise Exception("ffmpeg no está disponible en el sistema para procesar el archivo de audio.")

        with tempfile.TemporaryDirectory() as tmpdir:
            chunk_template = os.path.join(tmpdir, "chunk_%03d.wav")
            cmd = [
                ffmpeg_bin, "-y",
                "-i", file_path,
                "-f", "segment",
                "-segment_time", self.CHUNK_DURATION_SECONDS,
                "-c:a", "pcm_s16le",
                "-ar", "16000",
                "-ac", "1",
                chunk_template
            ]

            proc = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE)
            if proc.returncode != 0:
                raise Exception(f"ffmpeg no pudo decodificar el audio: {proc.stderr.decode('utf-8', errors='ignore')[:300]}")

            chunk_files = sorted(glob.glob(os.path.join(tmpdir, "chunk_*.wav")))
            if not chunk_files:
                raise Exception("No se pudieron generar segmentos de audio para transcribir.")

            recognizer = sr.Recognizer()
            transcribed_chunks = []

            for idx, chunk_file in enumerate(chunk_files):
                try:
                    with sr.AudioFile(chunk_file) as source:
                        audio_data = recognizer.record(source)
                    text = recognizer.recognize_google(audio_data, language="es-ES").strip()
                    if text:
                        transcribed_chunks.append(text)
                except sr.UnknownValueError:
                    continue
                except sr.RequestError as e:
                    sys.stderr.write(f"[STT Error] API error en segmento {idx+1}: {e}\n")
                    continue
                except Exception as e:
                    sys.stderr.write(f"[STT Error] Fallo en segmento {idx+1}: {e}\n")
                    continue

            if not transcribed_chunks:
                raise Exception("No se detectó voz o habla comprensible en el archivo de audio.")

            grouped_sections = []
            current_section = []
            for chunk in transcribed_chunks:
                formatted = chunk[0].upper() + chunk[1:]
                if not formatted.endswith(('.', '!', '?')):
                    formatted += '.'
                current_section.append(formatted)
                if len(current_section) >= 3:
                    grouped_sections.append(" ".join(current_section))
                    current_section = []

            if current_section:
                grouped_sections.append(" ".join(current_section))

            pages_text = [TextNormalizer.clean(s) for s in grouped_sections if TextNormalizer.clean(s)]
            clean_title = os.path.splitext(os.path.basename(file_path))[0]
            clean_title = re.sub(r'[_\-]+', ' ', clean_title).strip()
            clean_title = clean_title.title() if clean_title else "Transcripción de Audio"
            return ExtractionResult(title=clean_title, author="Transcripción de Audio (STT)", pages_text=pages_text, ocr_used=False)


# ──────────────────────────────────────────────────────────────────────────────
# 4. NLP PROCESSOR (Single Responsibility: Summarization & Chapter Chunking)
# ──────────────────────────────────────────────────────────────────────────────

class NlpProcessor:
    """Extractive NLP summarization and intelligent chapter track segmentation."""

    STOPWORDS = frozenset({
        'de', 'la', 'que', 'el', 'en', 'y', 'a', 'los', 'del', 'se', 'las', 'por', 'un', 'para', 'con', 'no',
        'una', 'su', 'al', 'lo', 'como', 'más', 'pero', 'sus', 'le', 'ya', 'o', 'este', 'sí', 'porque', 'esta',
        'entre', 'cuando', 'muy', 'sin', 'sobre', 'también', 'me', 'hasta', 'hay', 'donde', 'quien', 'desde',
        'todo', 'nos', 'durante', 'todos', 'uno', 'les', 'ni', 'contra', 'otros', 'ese', 'eso', 'ante', 'ellos',
        'estos', 'estas', 'ser', 'es', 'son', 'fue', 'era', 'sido', 'tener', 'tiene', 'tienen', 'hacer', 'hace',
        'the', 'of', 'and', 'to', 'in', 'is', 'for', 'that', 'with', 'on', 'as', 'at', 'this', 'by', 'an', 'be'
    })

    CHAPTER_REGEX = re.compile(
        r'(?:\n|^)(Capítulo|Capitulo|Chapter|Sección|Seccion|Módulo|Modulo|Tema|Parte|Unidad)\s+'
        r'([0-9IVXLCDM]+|[A-Z])[\.\:\s\-—]+([^\n]{2,80})',
        re.IGNORECASE
    )

    WORD_PATTERN = re.compile(r'\b[a-zA-ZáéíóúÁÉÍÓÚñÑ]{3,}\b')

    @classmethod
    def generate_summary(cls, text: str, max_sentences: int = 3, max_words: int = 140) -> str:
        if not text or len(text.strip()) < 50:
            return text.strip()

        raw_sentences = re.split(r'(?<=[.!?])\s+', text)
        sentences = [
            s.strip() for s in raw_sentences
            if len(s.strip()) > 25 and not s.strip().startswith(('http', 'www', '#', 'Capítulo', 'Chapter', 'Parte', 'Sección'))
        ]

        if not sentences:
            words = text.split()
            return ' '.join(words[:max_words])

        if len(sentences) <= max_sentences:
            return ' '.join(sentences)[:700]

        words = [w.lower() for w in cls.WORD_PATTERN.findall(text)]
        word_freq = {}
        for w in words:
            if w not in cls.STOPWORDS:
                word_freq[w] = word_freq.get(w, 0) + 1

        if not word_freq:
            return ' '.join(sentences[:max_sentences])

        scored_sentences = []
        for idx, sentence in enumerate(sentences):
            s_words = [w.lower() for w in cls.WORD_PATTERN.findall(sentence)]
            if not s_words:
                continue

            score = sum(word_freq.get(w, 0) for w in s_words) / len(s_words)
            position_ratio = idx / len(sentences)
            if position_ratio < 0.20:
                score *= 1.45
            elif position_ratio > 0.80:
                score *= 1.20

            word_count = len(s_words)
            if 10 <= word_count <= 35:
                score *= 1.25
            elif word_count < 8:
                score *= 0.60

            scored_sentences.append((score, idx, sentence))

        scored_sentences.sort(key=lambda x: x[0], reverse=True)
        top_candidates = sorted(scored_sentences[:max_sentences * 2], key=lambda x: x[1])

        summary_sentences = []
        current_words = 0
        for _, _, sentence in top_candidates:
            s_words = len(sentence.split())
            if current_words + s_words > max_words and summary_sentences:
                break
            summary_sentences.append(sentence)
            current_words += s_words
            if len(summary_sentences) >= max_sentences:
                break

        return ' '.join(summary_sentences).strip()

    @classmethod
    def segment_chapters(cls, full_text: str, default_title: str = "Capítulo") -> List[Dict]:
        matches = list(cls.CHAPTER_REGEX.finditer(full_text))
        chapters = []

        if len(matches) >= 2:
            for i, match in enumerate(matches):
                chap_type = match.group(1).capitalize()
                chap_num = match.group(2)
                chap_sub = match.group(3).strip()
                chap_title = f"{chap_type} {chap_num}: {chap_sub}"

                start_pos = match.start()
                end_pos = matches[i + 1].start() if i + 1 < len(matches) else len(full_text)

                chunk_text = full_text[start_pos:end_pos].strip()
                words = chunk_text.split()
                if len(words) >= 5:
                    chapters.append({
                        "chapter_number": len(chapters) + 1,
                        "title": chap_title[:100],
                        "text": chunk_text,
                        "word_count": len(words)
                    })

        if not chapters:
            paragraphs = full_text.split('\n\n')
            current_chunk = []
            current_words = 0
            target_words = 750
            chunk_idx = 1

            for p in paragraphs:
                p = p.strip()
                if not p:
                    continue
                p_words = len(p.split())
                current_chunk.append(p)
                current_words += p_words

                if current_words >= target_words:
                    chunk_text = '\n\n'.join(current_chunk)
                    preview = p[:60].replace('\n', ' ')
                    chap_title = f"Parte {chunk_idx}: {preview}..." if len(preview) > 10 else f"Parte {chunk_idx}"
                    chapters.append({
                        "chapter_number": chunk_idx,
                        "title": chap_title,
                        "text": chunk_text,
                        "word_count": current_words
                    })
                    current_chunk = []
                    current_words = 0
                    chunk_idx += 1

            if current_chunk:
                chunk_text = '\n\n'.join(current_chunk)
                preview = current_chunk[0][:60].replace('\n', ' ')
                chap_title = f"Parte {chunk_idx}: {preview}..." if len(preview) > 10 else f"Parte {chunk_idx}"
                chapters.append({
                    "chapter_number": chunk_idx,
                    "title": chap_title,
                    "text": chunk_text,
                    "word_count": current_words
                })

        return chapters


# ──────────────────────────────────────────────────────────────────────────────
# 5. ORCHESTRATION ENGINE (Strategy Selector & Pipeline Runner)
# ──────────────────────────────────────────────────────────────────────────────

class UniversalExtractionEngine:
    """Orchestrates document extraction by delegating to specific format strategies and running NLP pipelines."""

    STRATEGY_MAP: Dict[str, Type[BaseExtractor]] = {
        '.pdf': PdfExtractor,
        '.docx': DocxExtractor,
        '.doc': DocExtractor,
        '.txt': PlainTextExtractor,
        '.md': PlainTextExtractor,
        '.markdown': PlainTextExtractor,
        '.png': ImageOcrExtractor,
        '.jpg': ImageOcrExtractor,
        '.jpeg': ImageOcrExtractor,
        '.webp': ImageOcrExtractor,
        '.bmp': ImageOcrExtractor,
        '.mp3': AudioTranscriptionExtractor,
        '.wav': AudioTranscriptionExtractor,
        '.m4a': AudioTranscriptionExtractor,
        '.ogg': AudioTranscriptionExtractor,
        '.aac': AudioTranscriptionExtractor,
        '.flac': AudioTranscriptionExtractor,
    }

    @staticmethod
    def _detect_strategy_by_magic(file_path: str) -> Optional[Type[BaseExtractor]]:
        try:
            with open(file_path, 'rb') as f:
                header = f.read(32)
            if header.startswith(b'%PDF'):
                return PdfExtractor
            if (
                header.startswith(b'\x89PNG\r\n\x1a\n') or
                header.startswith(b'\xff\xd8\xff') or
                header.startswith(b'BM') or
                (header.startswith(b'RIFF') and b'WEBP' in header) or
                header.startswith(b'GIF87a') or
                header.startswith(b'GIF89a')
            ):
                return ImageOcrExtractor
            if header.startswith(b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1'):
                return DocExtractor
            if header.startswith(b'PK\x03\x04'):
                return DocxExtractor
            try:
                from PIL import Image
                with Image.open(file_path) as img:
                    img.verify()
                return ImageOcrExtractor
            except Exception:
                pass
        except Exception:
            pass
        return None

    def process(self, file_path: str) -> Dict:
        if not os.path.exists(file_path):
            return {"success": False, "error": f"Archivo no encontrado: {file_path}"}

        ext = os.path.splitext(file_path)[1].lower()
        extractor_class = self.STRATEGY_MAP.get(ext)

        if not extractor_class or ext in ('', '.tmp'):
            detected = self._detect_strategy_by_magic(file_path)
            if detected:
                extractor_class = detected

        if not extractor_class:
            return {"success": False, "error": f"Formato de archivo no soportado: {ext or 'sin extensión'}"}

        try:
            strategy = extractor_class()
            result: ExtractionResult = strategy.extract(file_path)
        except Exception as e:
            return {"success": False, "error": str(e)}

        total_text = "\n\n".join(result.pages_text).strip()
        total_words = len(total_text.split()) if total_text else 0

        if total_words < 2:
            return {
                "success": False,
                "error": (
                    "El documento no contiene texto legible ni voz reconocible. "
                    "Si es una imagen o PDF escaneado, asegúrate de que el texto sea claro y que Tesseract OCR esté disponible."
                )
            }

        title = result.title
        if not title:
            base_name = os.path.splitext(os.path.basename(file_path))[0]
            title = base_name.replace('_', ' ').replace('-', ' ').title()

        chapters = NlpProcessor.segment_chapters(total_text, default_title="Capítulo")
        summary = NlpProcessor.generate_summary(total_text, max_sentences=3, max_words=140)

        return {
            "success": True,
            "title": title,
            "author": result.author or "Autor Desconocido",
            "summary": summary,
            "total_pages": len(result.pages_text),
            "total_words": total_words,
            "total_chapters": len(chapters),
            "ocr_used": result.ocr_used,
            "file_type": ext.lstrip('.'),
            "chapters": chapters
        }


# ──────────────────────────────────────────────────────────────────────────────
# 6. CLI ENTRYPOINT
# ──────────────────────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(description="AudioLab MZ - Universal Document Extraction Engine v2.5")
    parser.add_argument("file_path", help="Path to input document, image or audio file")
    parser.add_argument("--output", help="Optional output JSON path", default=None)
    args = parser.parse_args()

    engine = UniversalExtractionEngine()
    payload = engine.process(args.file_path)

    json_output = json.dumps(payload, ensure_ascii=False, indent=2)

    if args.output:
        with open(args.output, "w", encoding="utf-8") as f:
            f.write(json_output)

    print(json_output)

    if not payload.get("success"):
        # If the file didn't exist or unsupported format, exit 1; if no readable words, exit 0
        if "no contiene texto legible" in payload.get("error", ""):
            sys.exit(0)
        sys.exit(1)


if __name__ == "__main__":
    main()
